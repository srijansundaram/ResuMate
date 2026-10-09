import io
import pdfplumber
import docx
import pytesseract
from PIL import Image


def extract_text(file_bytes: bytes, filename: str) -> tuple[str, dict]:
    """Returns (text, layout_info). layout_info feeds the ATS checker."""
    ext = filename.lower().rsplit(".", 1)[-1]
    info = {"pages": 1, "tables": 0, "images": 0, "ocr_used": False}
    text = ""

    if ext == "pdf":
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            info["pages"] = len(pdf.pages)
            for p in pdf.pages:
                text += (p.extract_text() or "") + "\n"
                info["tables"] += len(p.find_tables())
                info["images"] += len(p.images)
        if len(text.strip()) < 50:  # scanned PDF -> OCR
            from pdf2image import convert_from_bytes
            info["ocr_used"] = True
            for img in convert_from_bytes(file_bytes):
                text += pytesseract.image_to_string(img) + "\n"

    elif ext == "docx":
        d = docx.Document(io.BytesIO(file_bytes))
        text = "\n".join(p.text for p in d.paragraphs)
        info["tables"] = len(d.tables)
        info["images"] = len(d.inline_shapes)
        for t in d.tables:
            for row in t.rows:
                text += "\n" + " ".join(c.text for c in row.cells)

    elif ext in ("png", "jpg", "jpeg"):
        info["ocr_used"] = True
        text = pytesseract.image_to_string(Image.open(io.BytesIO(file_bytes)))

    else:  # txt
        text = file_bytes.decode("utf-8", errors="ignore")

    return text, info